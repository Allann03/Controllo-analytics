"""
Router de Classificação Contábil por Plano de Contas
=====================================================
CRUD do plano de contas por empresa, regras de classificação,
cadastro de clientes/fornecedores, mapeamento banco→conta,
e endpoint de classificação em lote de transações.
"""

import io
import uuid
from datetime import datetime, timezone
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from data.database.config import get_db
from data.database import models
from services.auth_utils import get_current_user as _get_user, resolve_empresa_or_403
from services.motor_classificacao import classificar_lote, deve_filtrar
from routers.auditoria import registrar_auditoria

router = APIRouter(prefix="/api/classificacao", tags=["classificacao"])


# ──────────────────────────────────────────────────────────────────
#  Helpers
# ──────────────────────────────────────────────────────────────────

def _agora():
    return datetime.now(timezone.utc).isoformat()


# ──────────────────────────────────────────────────────────────────
#  Schemas
# ──────────────────────────────────────────────────────────────────

class ContaPlanoCreate(BaseModel):
    codigo: str
    descricao: str
    tipo: str          # "sintetica" | "analitica"
    natureza: str      # "ativo" | "passivo" | "receita" | "despesa"
    codigo_pai: Optional[str] = ""

class ContaPlanoUpdate(BaseModel):
    descricao: Optional[str] = None
    tipo: Optional[str] = None
    natureza: Optional[str] = None
    codigo_pai: Optional[str] = None

class RegraCreate(BaseModel):
    padrao: str
    tipo_transacao: Optional[str] = "ambos"
    conta_debito_codigo: Optional[str] = ""
    conta_credito_codigo: Optional[str] = ""
    prioridade: Optional[int] = 0

class RegraUpdate(BaseModel):
    padrao: Optional[str] = None
    tipo_transacao: Optional[str] = None
    conta_debito_codigo: Optional[str] = None
    conta_credito_codigo: Optional[str] = None
    prioridade: Optional[int] = None
    ativo: Optional[bool] = None

class ContaBancoCreate(BaseModel):
    nome_banco: str
    conta_codigo: str
    conta_descricao: Optional[str] = ""

class ClienteFornecedorCreate(BaseModel):
    nome: str
    documento: Optional[str] = ""
    tipo: str          # "cliente" | "fornecedor"
    conta_codigo: Optional[str] = ""
    conta_descricao: Optional[str] = ""

class PlanoContasImportLote(BaseModel):
    contas: List[ContaPlanoCreate]


# ──────────────────────────────────────────────────────────────────
#  PLANO DE CONTAS DA EMPRESA — CRUD
# ──────────────────────────────────────────────────────────────────

@router.get("/plano/{empresa_id}")
def listar_plano_empresa(
    empresa_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    contas = (
        db.query(models.PlanoContasEmpresa)
        .filter(models.PlanoContasEmpresa.empresa_id == empresa_id)
        .order_by(models.PlanoContasEmpresa.ordem, models.PlanoContasEmpresa.codigo)
        .all()
    )
    return [
        {
            "id": c.id,
            "codigo": c.codigo,
            "descricao": c.descricao,
            "tipo": c.tipo,
            "natureza": c.natureza,
            "codigo_pai": c.codigo_pai,
            "ordem": c.ordem,
        }
        for c in contas
    ]


@router.post("/plano/{empresa_id}", status_code=201)
def criar_conta_plano(
    empresa_id: int,
    body: ContaPlanoCreate,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    existente = db.query(models.PlanoContasEmpresa).filter(
        models.PlanoContasEmpresa.empresa_id == empresa_id,
        models.PlanoContasEmpresa.codigo == body.codigo.strip(),
    ).first()
    if existente:
        raise HTTPException(status_code=400, detail=f"Conta {body.codigo} já existe para esta empresa.")

    agora = _agora()
    conta = models.PlanoContasEmpresa(
        empresa_id=empresa_id,
        codigo=body.codigo.strip(),
        descricao=body.descricao.strip(),
        tipo=body.tipo.strip().lower(),
        natureza=body.natureza.strip().lower(),
        codigo_pai=(body.codigo_pai or "").strip(),
        ordem=0,
        criado_em=agora,
        atualizado_em=agora,
    )
    db.add(conta)
    db.commit()
    db.refresh(conta)
    return {"id": conta.id, "codigo": conta.codigo, "descricao": conta.descricao}


@router.post("/plano/{empresa_id}/lote", status_code=201)
def importar_plano_lote(
    empresa_id: int,
    body: PlanoContasImportLote,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    """Importa várias contas do plano de uma vez (para carga inicial)."""
    resolve_empresa_or_403(db, empresa_id, user)
    agora = _agora()
    criadas = 0
    ignoradas = 0
    for c in body.contas:
        existente = db.query(models.PlanoContasEmpresa).filter(
            models.PlanoContasEmpresa.empresa_id == empresa_id,
            models.PlanoContasEmpresa.codigo == c.codigo.strip(),
        ).first()
        if existente:
            ignoradas += 1
            continue
        db.add(models.PlanoContasEmpresa(
            empresa_id=empresa_id,
            codigo=c.codigo.strip(),
            descricao=c.descricao.strip(),
            tipo=c.tipo.strip().lower(),
            natureza=c.natureza.strip().lower(),
            codigo_pai=(c.codigo_pai or "").strip(),
            ordem=0,
            criado_em=agora,
            atualizado_em=agora,
        ))
        criadas += 1
    db.commit()
    return {"criadas": criadas, "ignoradas": ignoradas}


@router.patch("/plano/{empresa_id}/{conta_id}")
def atualizar_conta_plano(
    empresa_id: int,
    conta_id: int,
    body: ContaPlanoUpdate,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    conta = db.query(models.PlanoContasEmpresa).filter(
        models.PlanoContasEmpresa.id == conta_id,
        models.PlanoContasEmpresa.empresa_id == empresa_id,
    ).first()
    if not conta:
        raise HTTPException(status_code=404, detail="Conta não encontrada.")
    if body.descricao is not None:
        conta.descricao = body.descricao.strip()
    if body.tipo is not None:
        conta.tipo = body.tipo.strip().lower()
    if body.natureza is not None:
        conta.natureza = body.natureza.strip().lower()
    if body.codigo_pai is not None:
        conta.codigo_pai = body.codigo_pai.strip()
    conta.atualizado_em = _agora()
    db.commit()
    return {"mensagem": "Conta atualizada."}


@router.delete("/plano/{empresa_id}/{conta_id}")
def remover_conta_plano(
    empresa_id: int,
    conta_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    conta = db.query(models.PlanoContasEmpresa).filter(
        models.PlanoContasEmpresa.id == conta_id,
        models.PlanoContasEmpresa.empresa_id == empresa_id,
    ).first()
    if not conta:
        raise HTTPException(status_code=404, detail="Conta não encontrada.")
    db.delete(conta)
    db.commit()
    return {"mensagem": "Conta removida."}


# ──────────────────────────────────────────────────────────────────
#  REGRAS DE CLASSIFICAÇÃO — CRUD
# ──────────────────────────────────────────────────────────────────

@router.get("/regras/{empresa_id}")
def listar_regras(
    empresa_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    regras = (
        db.query(models.RegraClassificacao)
        .filter(models.RegraClassificacao.empresa_id == empresa_id)
        .order_by(models.RegraClassificacao.prioridade.desc())
        .all()
    )
    # Busca descrições do plano de contas
    plano = {
        c.codigo: c.descricao
        for c in db.query(models.PlanoContasEmpresa)
        .filter(models.PlanoContasEmpresa.empresa_id == empresa_id)
        .all()
    }
    return [
        {
            "id": r.id,
            "padrao": r.padrao,
            "tipo_transacao": r.tipo_transacao,
            "conta_debito_codigo": r.conta_debito_codigo,
            "conta_debito_descricao": plano.get(r.conta_debito_codigo, ""),
            "conta_credito_codigo": r.conta_credito_codigo,
            "conta_credito_descricao": plano.get(r.conta_credito_codigo, ""),
            "prioridade": r.prioridade,
            "ativo": r.ativo,
        }
        for r in regras
    ]


@router.post("/regras/{empresa_id}", status_code=201)
def criar_regra(
    empresa_id: int,
    body: RegraCreate,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    agora = _agora()
    regra = models.RegraClassificacao(
        empresa_id=empresa_id,
        padrao=body.padrao.strip(),
        tipo_transacao=(body.tipo_transacao or "ambos").strip().lower(),
        conta_debito_codigo=(body.conta_debito_codigo or "").strip(),
        conta_credito_codigo=(body.conta_credito_codigo or "").strip(),
        prioridade=body.prioridade or 0,
        ativo=True,
        criado_em=agora,
        atualizado_em=agora,
    )
    db.add(regra)
    db.commit()
    db.refresh(regra)
    registrar_auditoria(db, "criar_regra_classificacao", usuario_id=user.id, usuario_nome=user.nome,
                        escritorio_id=user.escritorio_id, recurso="regra", recurso_id=regra.id,
                        detalhes={"empresa_id": empresa_id, "padrao": regra.padrao})
    return {"id": regra.id, "padrao": regra.padrao}


@router.patch("/regras/{empresa_id}/{regra_id}")
def atualizar_regra(
    empresa_id: int,
    regra_id: int,
    body: RegraUpdate,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    regra = db.query(models.RegraClassificacao).filter(
        models.RegraClassificacao.id == regra_id,
        models.RegraClassificacao.empresa_id == empresa_id,
    ).first()
    if not regra:
        raise HTTPException(status_code=404, detail="Regra não encontrada.")
    if body.padrao is not None:
        regra.padrao = body.padrao.strip()
    if body.tipo_transacao is not None:
        regra.tipo_transacao = body.tipo_transacao.strip().lower()
    if body.conta_debito_codigo is not None:
        regra.conta_debito_codigo = body.conta_debito_codigo.strip()
    if body.conta_credito_codigo is not None:
        regra.conta_credito_codigo = body.conta_credito_codigo.strip()
    if body.prioridade is not None:
        regra.prioridade = body.prioridade
    if body.ativo is not None:
        regra.ativo = body.ativo
    regra.atualizado_em = _agora()
    db.commit()
    return {"mensagem": "Regra atualizada."}


@router.delete("/regras/{empresa_id}/{regra_id}")
def remover_regra(
    empresa_id: int,
    regra_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    regra = db.query(models.RegraClassificacao).filter(
        models.RegraClassificacao.id == regra_id,
        models.RegraClassificacao.empresa_id == empresa_id,
    ).first()
    if not regra:
        raise HTTPException(status_code=404, detail="Regra não encontrada.")
    padrao_regra = regra.padrao
    db.delete(regra)
    db.commit()
    registrar_auditoria(db, "remover_regra_classificacao", usuario_id=user.id, usuario_nome=user.nome,
                        escritorio_id=user.escritorio_id, recurso="regra", recurso_id=regra_id,
                        detalhes={"empresa_id": empresa_id, "padrao": padrao_regra})
    return {"mensagem": "Regra removida."}


# ──────────────────────────────────────────────────────────────────
#  CONTA BANCO → CONTA ANALÍTICA
# ──────────────────────────────────────────────────────────────────

@router.get("/bancos/{empresa_id}")
def listar_contas_banco(
    empresa_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    contas = (
        db.query(models.ContaBancoEmpresa)
        .filter(models.ContaBancoEmpresa.empresa_id == empresa_id)
        .order_by(models.ContaBancoEmpresa.nome_banco)
        .all()
    )
    return [
        {
            "id": c.id,
            "nome_banco": c.nome_banco,
            "conta_codigo": c.conta_codigo,
            "conta_descricao": c.conta_descricao,
        }
        for c in contas
    ]


@router.post("/bancos/{empresa_id}", status_code=201)
def criar_conta_banco(
    empresa_id: int,
    body: ContaBancoCreate,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    existente = db.query(models.ContaBancoEmpresa).filter(
        models.ContaBancoEmpresa.empresa_id == empresa_id,
        models.ContaBancoEmpresa.nome_banco == body.nome_banco.strip(),
    ).first()
    if existente:
        raise HTTPException(status_code=400, detail=f"Banco '{body.nome_banco}' já cadastrado.")
    agora = _agora()
    conta = models.ContaBancoEmpresa(
        empresa_id=empresa_id,
        nome_banco=body.nome_banco.strip(),
        conta_codigo=body.conta_codigo.strip(),
        conta_descricao=(body.conta_descricao or "").strip(),
        criado_em=agora,
        atualizado_em=agora,
    )
    db.add(conta)
    db.commit()
    db.refresh(conta)
    return {"id": conta.id, "nome_banco": conta.nome_banco}


@router.delete("/bancos/{empresa_id}/{conta_id}")
def remover_conta_banco(
    empresa_id: int,
    conta_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    conta = db.query(models.ContaBancoEmpresa).filter(
        models.ContaBancoEmpresa.id == conta_id,
        models.ContaBancoEmpresa.empresa_id == empresa_id,
    ).first()
    if not conta:
        raise HTTPException(status_code=404, detail="Registro não encontrado.")
    db.delete(conta)
    db.commit()
    return {"mensagem": "Mapeamento de banco removido."}


# ──────────────────────────────────────────────────────────────────
#  CLIENTES / FORNECEDORES
# ──────────────────────────────────────────────────────────────────

@router.get("/cadastros/{empresa_id}")
def listar_cadastros(
    empresa_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    cadastros = (
        db.query(models.ClienteFornecedorEmpresa)
        .filter(models.ClienteFornecedorEmpresa.empresa_id == empresa_id)
        .order_by(models.ClienteFornecedorEmpresa.nome)
        .all()
    )
    return [
        {
            "id": c.id,
            "nome": c.nome,
            "documento": c.documento,
            "tipo": c.tipo,
            "conta_codigo": c.conta_codigo,
            "conta_descricao": c.conta_descricao,
        }
        for c in cadastros
    ]


@router.post("/cadastros/{empresa_id}", status_code=201)
def criar_cadastro(
    empresa_id: int,
    body: ClienteFornecedorCreate,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    agora = _agora()
    cad = models.ClienteFornecedorEmpresa(
        empresa_id=empresa_id,
        nome=body.nome.strip(),
        documento=(body.documento or "").strip(),
        tipo=body.tipo.strip().lower(),
        conta_codigo=(body.conta_codigo or "").strip(),
        conta_descricao=(body.conta_descricao or "").strip(),
        criado_em=agora,
        atualizado_em=agora,
    )
    db.add(cad)
    db.commit()
    db.refresh(cad)
    return {"id": cad.id, "nome": cad.nome}


@router.delete("/cadastros/{empresa_id}/{cad_id}")
def remover_cadastro(
    empresa_id: int,
    cad_id: int,
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    resolve_empresa_or_403(db, empresa_id, user)
    cad = db.query(models.ClienteFornecedorEmpresa).filter(
        models.ClienteFornecedorEmpresa.id == cad_id,
        models.ClienteFornecedorEmpresa.empresa_id == empresa_id,
    ).first()
    if not cad:
        raise HTTPException(status_code=404, detail="Cadastro não encontrado.")
    db.delete(cad)
    db.commit()
    return {"mensagem": "Cadastro removido."}


# ──────────────────────────────────────────────────────────────────
#  HELPERS — leitura do Excel de input e carga de dados
# ──────────────────────────────────────────────────────────────────

def _ler_excel_input(content: bytes) -> list:
    """Lê Excel do extrator e retorna lista de transações."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active

    header = [str(c.value or "").strip().lower() for c in ws[1]]
    col_map = {}
    for i, h in enumerate(header):
        if "data" in h:
            col_map["data"] = i
        elif "banco" in h:
            col_map["banco"] = i
        elif "categ" in h:
            col_map["categoria"] = i
        elif "descri" in h:
            col_map["descricao"] = i
        elif "tipo" in h:
            col_map["tipo"] = i
        elif "valor" in h:
            col_map["valor"] = i

    required = ["data", "descricao", "tipo", "valor"]
    missing = [r for r in required if r not in col_map]
    if missing:
        raise HTTPException(status_code=400, detail=f"Colunas não encontradas: {', '.join(missing)}")

    transacoes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        vals = list(row)
        data_val = vals[col_map["data"]] if col_map["data"] < len(vals) else ""
        banco_val = vals[col_map.get("banco", 0)] if "banco" in col_map and col_map["banco"] < len(vals) else ""
        desc_val = vals[col_map["descricao"]] if col_map["descricao"] < len(vals) else ""
        tipo_val = vals[col_map["tipo"]] if col_map["tipo"] < len(vals) else ""
        valor_val = vals[col_map["valor"]] if col_map["valor"] < len(vals) else 0

        if desc_val is None or str(desc_val).strip() == "":
            continue

        if hasattr(data_val, "strftime"):
            data_str = data_val.strftime("%d/%m/%Y")
        else:
            data_str = str(data_val or "")

        try:
            valor_num = float(str(valor_val).replace("R$", "").replace(".", "").replace(",", ".").strip())
        except (ValueError, TypeError):
            valor_num = 0.0

        transacoes.append({
            "data": data_str,
            "banco": str(banco_val or ""),
            "descricao": str(desc_val or ""),
            "tipo": str(tipo_val or ""),
            "valor": valor_num,
        })
    return transacoes


def _carregar_dados_classificacao(empresa_id: int, db: Session) -> tuple:
    """Carrega todos os dados necessários para classificação."""
    # Mapeamento banco → código reduzido
    conta_banco_map = {}
    for cb in db.query(models.ContaBancoEmpresa).filter(
        models.ContaBancoEmpresa.empresa_id == empresa_id
    ).all():
        conta_banco_map[cb.nome_banco] = cb.conta_codigo

    # Regras do usuário
    regras_usuario = []
    for r in db.query(models.RegraClassificacao).filter(
        models.RegraClassificacao.empresa_id == empresa_id,
        models.RegraClassificacao.ativo == True,
    ).order_by(models.RegraClassificacao.prioridade.desc()).all():
        regras_usuario.append({
            "padrao": r.padrao,
            "tipo_transacao": r.tipo_transacao,
            "conta_debito_codigo": r.conta_debito_codigo,
            "conta_credito_codigo": r.conta_credito_codigo,
        })

    # Cadastro de pessoas
    cadastros = []
    for c in db.query(models.ClienteFornecedorEmpresa).filter(
        models.ClienteFornecedorEmpresa.empresa_id == empresa_id
    ).all():
        cadastros.append({
            "nome": c.nome,
            "documento": c.documento,
            "tipo": c.tipo,
            "conta_codigo": c.conta_codigo,
        })

    # Plano de contas: código reduzido → descrição
    plano = {
        c.codigo: c.descricao
        for c in db.query(models.PlanoContasEmpresa)
        .filter(models.PlanoContasEmpresa.empresa_id == empresa_id)
        .all()
    }

    return conta_banco_map, regras_usuario, cadastros, plano


# ──────────────────────────────────────────────────────────────────
#  CLASSIFICAR EXCEL — gera Excel padrão Contmatic
# ──────────────────────────────────────────────────────────────────

@router.post("/classificar/{empresa_id}")
async def classificar_excel(
    empresa_id: int,
    arquivo: UploadFile = File(...),
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    """
    Recebe Excel do extrator, classifica e retorna Excel no padrão Contmatic:
    Data | Lançamento | Histórico | Descrição | Débito | Crédito | Valor (R$)
    """
    resolve_empresa_or_403(db, empresa_id, user)

    if not arquivo.filename or not arquivo.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Apenas arquivos .xlsx são aceitos.")

    content = await arquivo.read()
    if not content:
        raise HTTPException(status_code=400, detail="Arquivo vazio.")

    transacoes = _ler_excel_input(content)
    if not transacoes:
        raise HTTPException(status_code=400, detail="Nenhuma transação encontrada no Excel.")

    conta_banco_map, regras_usuario, cadastros, plano = _carregar_dados_classificacao(empresa_id, db)

    resultado = classificar_lote(
        transacoes=transacoes,
        conta_banco_map=conta_banco_map,
        regras_usuario=regras_usuario,
        cadastros=cadastros,
        plano=plano,
    )

    # Gera Excel padrão Contmatic (7 colunas)
    import openpyxl
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Contmatic"

    headers = ["Data", "Lançamento", "Histórico", "Descrição", "Débito", "Crédito", "Valor (R$)"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws_out.cell(row=1, column=col_idx, value=h)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")

    status_colors = {
        "regra_usuario": "E8F5E9",
        "cadastro": "E3F2FD",
        "automatico": "FFFDE7",
        "pendente": "FFEBEE",
    }
    for row_idx, tx in enumerate(resultado, 2):
        ws_out.cell(row=row_idx, column=1, value=tx["data"])
        ws_out.cell(row=row_idx, column=2, value=tx["lancamento"])
        ws_out.cell(row=row_idx, column=3, value=tx["historico"])
        ws_out.cell(row=row_idx, column=4, value=tx["descricao"])
        ws_out.cell(row=row_idx, column=5, value=tx["debito"])
        ws_out.cell(row=row_idx, column=6, value=tx["credito"])
        ws_out.cell(row=row_idx, column=7, value=tx["valor"])

        bg = status_colors.get(tx["status"], "FFFFFF")
        fill = openpyxl.styles.PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        for col in range(1, 8):
            ws_out.cell(row=row_idx, column=col).fill = fill

    # Formato
    for row in ws_out.iter_rows(min_row=2, min_col=7, max_col=7):
        for cell in row:
            cell.number_format = '#,##0.00'

    widths = [12, 12, 10, 55, 10, 10, 15]
    for i, w in enumerate(widths, 1):
        ws_out.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb_out.save(output)
    output.seek(0)

    filename = f"contmatic_{empresa_id}_{uuid.uuid4().hex[:8]}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/preview/{empresa_id}")
async def preview_classificacao(
    empresa_id: int,
    arquivo: UploadFile = File(...),
    user: models.Usuario = Depends(_get_user),
    db: Session = Depends(get_db),
):
    """Preview da classificação — retorna JSON no formato Contmatic."""
    resolve_empresa_or_403(db, empresa_id, user)

    if not arquivo.filename or not arquivo.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Apenas arquivos .xlsx são aceitos.")

    content = await arquivo.read()
    if not content:
        raise HTTPException(status_code=400, detail="Arquivo vazio.")

    transacoes = _ler_excel_input(content)
    if not transacoes:
        raise HTTPException(status_code=400, detail="Nenhuma transação encontrada.")

    conta_banco_map, regras_usuario, cadastros, plano = _carregar_dados_classificacao(empresa_id, db)

    resultado = classificar_lote(
        transacoes=transacoes,
        conta_banco_map=conta_banco_map,
        regras_usuario=regras_usuario,
        cadastros=cadastros,
        plano=plano,
    )

    stats = {
        "total": len(resultado),
        "regra_usuario": sum(1 for t in resultado if t["status"] == "regra_usuario"),
        "cadastro": sum(1 for t in resultado if t["status"] == "cadastro"),
        "automatico": sum(1 for t in resultado if t["status"] == "automatico"),
        "pendente": sum(1 for t in resultado if t["status"] == "pendente"),
        "filtradas": len(transacoes) - len(resultado),
    }

    return {"transacoes": resultado, "stats": stats}

