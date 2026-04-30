"""
Testes do gerador Excel básico — spec da Sessão 22 (Allan).

Layout: 7 colunas (Lançamento, Data, Débito, Crédito, Valor, Histórico, Complemento),
1 aba, sem totalizadores, valor sempre positivo.
"""
import os
import sys
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from services.gerador_excel import gerar_excel  # noqa: E402


CABECALHOS_ESPERADOS = [
    "Lançamento", "Data", "Débito", "Crédito", "Valor", "Histórico", "Complemento"
]


def _tx(data, descricao, valor, tipo="entrada", banco="nubank", categoria="Receita"):
    return {
        "data": data,
        "descricao": descricao,
        "valor": valor,
        "tipo": tipo,
        "banco": banco,
        "categoria": categoria,
        "raw": "",
    }


@pytest.fixture
def excel_path(tmp_path):
    return str(tmp_path / "saida.xlsx")


def test_cabecalhos_e_ordem_de_colunas(excel_path):
    txs = [_tx("01/01/2026", "PIX recebido", Decimal("100.00"))]
    gerar_excel(txs, excel_path)

    wb = load_workbook(excel_path)
    assert wb.sheetnames == ["Lançamentos"], "deve ter exatamente 1 aba"
    ws = wb.active
    cabecalhos = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert cabecalhos == CABECALHOS_ESPERADOS
    # Cabeçalho em negrito
    assert ws.cell(row=1, column=1).font.bold is True


def test_colunas_debito_credito_historico_vazias(excel_path):
    txs = [
        _tx("01/01/2026", "Receita A", Decimal("50.00"), tipo="entrada"),
        _tx("02/01/2026", "Despesa B", Decimal("-30.00"), tipo="saida"),
    ]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    for row in (2, 3):
        assert ws.cell(row=row, column=3).value in (None, ""), "Débito (C) deve ficar vazia"
        assert ws.cell(row=row, column=4).value in (None, ""), "Crédito (D) deve ficar vazia"
        assert ws.cell(row=row, column=6).value in (None, ""), "Histórico (F) deve ficar vazia"


def test_valor_sempre_positivo_inclusive_para_saida_negativa(excel_path):
    """Pega regressão se alguém remover o abs() — parsers como Itaú emitem Decimal negativo para saída."""
    txs = [
        _tx("01/01/2026", "Crédito",  Decimal("100.00"),  tipo="entrada"),
        _tx("02/01/2026", "Débito",   Decimal("-300.00"), tipo="saida"),
        _tx("03/01/2026", "Boleto",   Decimal("-1234.56"), tipo="saida"),
    ]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    valores = [ws.cell(row=r, column=5).value for r in (2, 3, 4)]
    assert valores == [100.00, 300.00, 1234.56]
    for r in (2, 3, 4):
        assert ws.cell(row=r, column=5).number_format == "#,##0.00"


def test_lancamento_numeracao_crescente(excel_path):
    txs = [_tx("01/01/2026", f"tx {i}", Decimal("10.00")) for i in range(5)]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    lancamentos = [ws.cell(row=r, column=1).value for r in range(2, 7)]
    assert lancamentos == [1, 2, 3, 4, 5]


def test_data_como_tipo_excel_e_complemento_descricao(excel_path):
    txs = [_tx("15/03/2026", "Pagamento NF 123", Decimal("200.00"))]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    data_cell = ws.cell(row=2, column=2)
    assert isinstance(data_cell.value, (date, datetime)), "Data deve ser tipo date/datetime"
    assert data_cell.value.day == 15 and data_cell.value.month == 3 and data_cell.value.year == 2026
    assert data_cell.number_format == "DD/MM/YYYY"
    assert ws.cell(row=2, column=7).value == "Pagamento NF 123"


def test_lista_vazia_gera_excel_com_so_cabecalho(excel_path):
    """Caso de borda: PDF que retorna 0 transações ainda precisa devolver Excel."""
    gerar_excel([], excel_path)

    wb = load_workbook(excel_path)
    assert wb.sheetnames == ["Lançamentos"]
    ws = wb.active
    cabecalhos = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert cabecalhos == CABECALHOS_ESPERADOS
    # Linha 2 não deve ter dados
    assert ws.cell(row=2, column=1).value is None


def test_filtra_tipo_posicao(excel_path):
    """Posições XP (carteira de investimento) NÃO são lançamentos contábeis — filtradas."""
    txs = [
        _tx("01/01/2026", "PIX X",          Decimal("100.00"), tipo="entrada"),
        _tx("01/01/2026", "PETR4 100 ações", Decimal("5000.00"), tipo="posicao", banco="xp_posicao", categoria="Posição de Investimentos"),
        _tx("02/01/2026", "Boleto",         Decimal("-50.00"),  tipo="saida"),
    ]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    # 2 linhas de dados, não 3
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=4, column=1).value is None
    # PETR4 não pode estar em nenhuma linha
    complementos = [ws.cell(row=r, column=7).value for r in range(2, 5)]
    assert all("PETR4" not in (c or "") for c in complementos)


def test_sem_abas_extras_sem_totalizadores(excel_path):
    txs = [
        _tx("01/01/2026", "A", Decimal("100.00"), tipo="entrada"),
        _tx("02/01/2026", "B", Decimal("-50.00"), tipo="saida"),
    ]
    gerar_excel(txs, excel_path, saldo_inicial=1000.00)  # parâmetro aceito mas ignorado

    wb = load_workbook(excel_path)
    assert wb.sheetnames == ["Lançamentos"], "não pode ter aba Resumo nem Plano de Contas"
    ws = wb.active
    # Última linha de dados é 3 (cab + 2 txs). Linha 4 não pode ter "TOTAIS" / "Saldo"
    for row in range(4, 10):
        for col in range(1, 8):
            v = ws.cell(row=row, column=col).value
            assert v is None, f"linha {row} col {col} tem valor {v!r} — não deveria haver totalizador"


# ---------------------------------------------------------------------------
# Extensão S22: indicador visual de cor (cabeçalho azul + linhas verde/rosa).
# ---------------------------------------------------------------------------

# openpyxl normaliza cores para "AARRGGBB" (8 chars com canal alpha).
COR_HEADER_BG = "FF1F4E79"
COR_HEADER_FG = "FFFFFFFF"
COR_ENTRADA_BG = "FFE2EFDA"
COR_SAIDA_BG = "FFFCE4D6"


def _cor_fill(cell):
    """Lê a cor de fundo de uma célula como string AARRGGBB (ou None se sem fill)."""
    if cell.fill is None or cell.fill.fgColor is None:
        return None
    return cell.fill.fgColor.rgb


def test_cabecalho_tem_fundo_azul_e_texto_branco(excel_path):
    txs = [_tx("01/01/2026", "X", Decimal("10.00"), tipo="entrada")]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    for col in range(1, 8):
        cell = ws.cell(row=1, column=col)
        assert _cor_fill(cell) == COR_HEADER_BG, f"col {col}: fundo cabeçalho deve ser azul {COR_HEADER_BG}"
        assert cell.font.color.rgb == COR_HEADER_FG, f"col {col}: texto cabeçalho deve ser branco"
        assert cell.font.bold is True, f"col {col}: cabeçalho deve permanecer em negrito"


def test_linha_entrada_tem_fundo_verde_em_todas_colunas(excel_path):
    txs = [
        _tx("01/01/2026", "Receita", Decimal("100.00"), tipo="entrada"),
        _tx("02/01/2026", "Despesa", Decimal("-50.00"), tipo="saida"),
    ]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    # Linha 2 = entrada → verde em A:G (incluindo C, D, F que são vazias)
    for col in range(1, 8):
        assert _cor_fill(ws.cell(row=2, column=col)) == COR_ENTRADA_BG, (
            f"linha entrada col {col}: fundo deve ser verde {COR_ENTRADA_BG}"
        )


def test_linha_saida_tem_fundo_rosa_em_todas_colunas(excel_path):
    txs = [
        _tx("01/01/2026", "Receita", Decimal("100.00"), tipo="entrada"),
        _tx("02/01/2026", "Despesa", Decimal("-50.00"), tipo="saida"),
    ]
    gerar_excel(txs, excel_path)

    ws = load_workbook(excel_path).active
    # Linha 3 = saida → rosa em A:G (incluindo C, D, F que são vazias)
    for col in range(1, 8):
        assert _cor_fill(ws.cell(row=3, column=col)) == COR_SAIDA_BG, (
            f"linha saida col {col}: fundo deve ser rosa {COR_SAIDA_BG}"
        )
