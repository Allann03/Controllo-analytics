"""
Gerador de Excel Contábil — Partida Dobrada (padrão Contmatic/Domínio/Prosoft)
==============================================================================

Recebe transações já extraídas do PDF e gera Excel com:
  - Aba 1: "Lançamentos Contábeis" (Data|Lançamento|Histórico|Descrição|Débito|Crédito|Valor)
  - Aba 2: "Resumo" (totais, conta do banco, estatísticas de classificação)
  - Aba 3: "Plano de Contas Utilizado" (contas que aparecem nos lançamentos)

Usa o motor_classificacao.py para classificar cada transação.
"""

from decimal import Decimal
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from services.motor_classificacao import classificar_lote


# ── Estilos ──────────────────────────────────────────────────────────

_COR_HEADER = "1E3A5F"
_COR_HEADER_FONT = "FFFFFF"
_COR_REGRA_USUARIO = "E8F5E9"   # verde claro
_COR_CADASTRO = "E3F2FD"        # azul claro
_COR_AUTOMATICO = "FFFDE7"      # amarelo claro
_COR_PENDENTE = "FFEBEE"        # vermelho claro
_COR_BORDA = "CCCCCC"

_STATUS_CORES = {
    "regra_usuario": _COR_REGRA_USUARIO,
    "cadastro": _COR_CADASTRO,
    "automatico": _COR_AUTOMATICO,
    "pendente": _COR_PENDENTE,
}

_STATUS_LABELS = {
    "regra_usuario": "Regra do usuário",
    "cadastro": "Cadastro cliente/fornecedor",
    "automatico": "Classificação automática",
    "pendente": "A Classificar (revisar)",
}


def _to_float(val) -> float:
    if isinstance(val, Decimal):
        return float(val)
    if val is None:
        return 0.0
    return float(val)


def _header_cell(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Arial", bold=True, color=_COR_HEADER_FONT, size=11)
    cell.fill = PatternFill(start_color=_COR_HEADER, end_color=_COR_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    return cell


def _section_header(ws, row, text, cols=7):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Arial", bold=True, size=11, color="1E3A5F")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)


def gerar_excel_contabil(
    transacoes: list,
    conta_banco_map: dict,
    regras_usuario: list,
    cadastros: list,
    plano: dict,
    banco_detectado: str = "",
    saldo_inicial: float = 0.0,
    saldo_final: float = None,
) -> bytes:
    """
    Gera Excel contábil com partida dobrada.

    Args:
        transacoes: lista de dicts com data, descricao, valor, tipo, banco, etc.
        conta_banco_map: dict {nome_banco: codigo_reduzido}
        regras_usuario: lista de regras customizadas
        cadastros: lista de cadastros (clientes/fornecedores)
        plano: dict {codigo: descricao} do plano de contas
        banco_detectado: nome do banco para exibição
        saldo_inicial: saldo inicial do extrato
        saldo_final: saldo final informado no extrato

    Returns:
        bytes do arquivo Excel
    """
    # ── Garantir que nomes de banco das transações estejam no map ──
    # (o campo 'banco' nas transações usa nome de exibição, não a chave do parser)
    _conta_default = next(iter(conta_banco_map.values()), "11200") if conta_banco_map else "11200"
    for t in transacoes:
        b = t.get("banco", "")
        if b and b not in conta_banco_map:
            conta_banco_map[b] = _conta_default

    # ── Classificar transações ───────────────────────────────────
    lancamentos = classificar_lote(
        transacoes=transacoes,
        conta_banco_map=conta_banco_map,
        regras_usuario=regras_usuario,
        cadastros=cadastros,
        plano=plano,
    )

    wb = Workbook()

    # ════════════════════════════════════════════════════════════
    # ABA 1: Lançamentos Contábeis
    # ════════════════════════════════════════════════════════════
    ws = wb.active
    ws.title = "Lançamentos Contábeis"

    headers = ["Data", "Lançamento", "Histórico", "Descrição", "Débito", "Crédito", "Valor (R$)"]
    for col_idx, h in enumerate(headers, 1):
        _header_cell(ws, 1, col_idx, h)

    borda = Border(
        left=Side(style="thin", color=_COR_BORDA),
        right=Side(style="thin", color=_COR_BORDA),
        top=Side(style="thin", color=_COR_BORDA),
        bottom=Side(style="thin", color=_COR_BORDA),
    )

    total_entradas = 0.0
    total_saidas = 0.0
    contas_usadas = {}  # codigo -> {descricao, count, total}

    for row_idx, tx in enumerate(lancamentos, 2):
        valor = _to_float(tx["valor"])

        ws.cell(row=row_idx, column=1, value=tx["data"])
        ws.cell(row=row_idx, column=2, value=tx["lancamento"])
        ws.cell(row=row_idx, column=3, value=tx["historico"])
        ws.cell(row=row_idx, column=4, value=tx["descricao"])
        ws.cell(row=row_idx, column=5, value=tx["debito"])
        ws.cell(row=row_idx, column=6, value=tx["credito"])
        cel_val = ws.cell(row=row_idx, column=7, value=valor)
        cel_val.number_format = '#,##0.00'

        # Cor de fundo por status de classificação
        bg = _STATUS_CORES.get(tx.get("status", ""), "FFFFFF")
        fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        for col in range(1, 8):
            c = ws.cell(row=row_idx, column=col)
            c.fill = fill
            c.border = borda
            if col in (5, 6):
                c.alignment = Alignment(horizontal="center")

        # Acumular totais
        # Determinar se é entrada ou saída baseado na posição da conta banco
        is_entrada = False
        for banco_nome, banco_codigo in conta_banco_map.items():
            if tx["debito"] == banco_codigo:
                is_entrada = True
                break

        if is_entrada:
            total_entradas += valor
        else:
            total_saidas += valor

        # Rastrear contas usadas
        for codigo in (tx["debito"], tx["credito"]):
            if codigo:
                if codigo not in contas_usadas:
                    desc = plano.get(codigo, "")
                    contas_usadas[codigo] = {"descricao": desc, "count": 0, "total": 0.0}
                contas_usadas[codigo]["count"] += 1
                contas_usadas[codigo]["total"] += valor

    # Larguras das colunas
    widths = [12, 12, 10, 55, 10, 10, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ════════════════════════════════════════════════════════════
    # ABA 2: Resumo
    # ════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Resumo")

    _section_header(ws2, 1, "Resumo do Extrato Contábil", 4)

    si = _to_float(saldo_inicial)
    sf_calc = si + total_entradas - total_saidas
    sf_info = _to_float(saldo_final) if saldo_final is not None else sf_calc

    # Conta do banco
    conta_banco_usado = ""
    for nome, codigo in conta_banco_map.items():
        if banco_detectado.lower() in nome.lower() or nome.lower() in banco_detectado.lower():
            conta_banco_usado = f"{codigo} ({nome})"
            break
    if not conta_banco_usado and conta_banco_map:
        primeiro = next(iter(conta_banco_map.items()))
        conta_banco_usado = f"{primeiro[1]} ({primeiro[0]})"

    # Estatísticas de classificação
    total_lanc = len(lancamentos)
    stats = {}
    for tx in lancamentos:
        s = tx.get("status", "pendente")
        stats[s] = stats.get(s, 0) + 1

    auto = stats.get("automatico", 0) + stats.get("regra_usuario", 0) + stats.get("cadastro", 0)
    pendente = stats.get("pendente", 0)
    pct_auto = round(auto / total_lanc * 100, 1) if total_lanc > 0 else 0

    resumo_dados = [
        ("Banco", banco_detectado or "Não identificado"),
        ("Conta contábil do banco", conta_banco_usado or "Não cadastrada"),
        ("", ""),
        ("Saldo Inicial", si),
        ("Total Entradas", total_entradas),
        ("Total Saídas", total_saidas),
        ("Saldo Final (calculado)", sf_calc),
        ("Saldo Final (informado)", sf_info),
        ("Conferência", "OK" if abs(sf_calc - sf_info) < 0.02 else f"DIVERGÊNCIA: R$ {sf_calc - sf_info:,.2f}"),
        ("", ""),
        ("Total de lançamentos", total_lanc),
        ("Classificados automaticamente", f"{auto} ({pct_auto}%)"),
        ("A Classificar (revisar)", f"{pendente} ({round(pendente/total_lanc*100, 1) if total_lanc else 0}%)"),
        ("", ""),
        ("Legenda de cores", ""),
    ]

    header_font = Font(name="Arial", bold=True, size=11, color="1E3A5F")
    for i, (label, value) in enumerate(resumo_dados, 3):
        c1 = ws2.cell(row=i, column=1, value=label)
        c2 = ws2.cell(row=i, column=2, value=value)
        c1.font = Font(name="Arial", bold=True, size=10)
        if isinstance(value, float):
            c2.number_format = '#,##0.00'

    # Legenda de cores
    row_legenda = 3 + len(resumo_dados)
    for status, label in _STATUS_LABELS.items():
        cor = _STATUS_CORES[status]
        c1 = ws2.cell(row=row_legenda, column=1, value=label)
        c1.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
        c1.font = Font(name="Arial", size=9)
        row_legenda += 1

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 30

    # ════════════════════════════════════════════════════════════
    # ABA 3: Plano de Contas Utilizado
    # ════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("Plano de Contas Utilizado")

    headers3 = ["Código", "Descrição", "Qtd Lançamentos", "Total R$"]
    for col_idx, h in enumerate(headers3, 1):
        _header_cell(ws3, 1, col_idx, h)

    for row_idx, (codigo, info) in enumerate(
        sorted(contas_usadas.items(), key=lambda x: x[0]), 2
    ):
        ws3.cell(row=row_idx, column=1, value=codigo)
        ws3.cell(row=row_idx, column=2, value=info["descricao"])
        ws3.cell(row=row_idx, column=3, value=info["count"])
        cel = ws3.cell(row=row_idx, column=4, value=info["total"])
        cel.number_format = '#,##0.00'

    ws3.column_dimensions["A"].width = 15
    ws3.column_dimensions["B"].width = 40
    ws3.column_dimensions["C"].width = 18
    ws3.column_dimensions["D"].width = 18

    # ── Salvar em bytes ──────────────────────────────────────────
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
