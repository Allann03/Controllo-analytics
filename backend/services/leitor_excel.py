from openpyxl import load_workbook
from io import BytesIO

def ler_planilhas_excel(arquivos_bytes: list) -> list:
    transacoes_consolidadas = []

    # Como o usuário pode subir vários meses de uma vez, fazemos um loop por todos os arquivos
    for arquivo_bytes in arquivos_bytes:
        # Carrega a planilha diretamente da memória (sem precisar salvar no computador do servidor)
        wb = load_workbook(filename=BytesIO(arquivo_bytes), data_only=True)
        
        # Garante que vai ler a aba certa (se existir "Transações", ele pega. Se não, pega a primeira que achar)
        nome_aba = "Transações" if "Transações" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[nome_aba]

        # Pula a linha 1 (que é o cabeçalho) e começa a ler da linha 2 em diante
        for row in ws.iter_rows(min_row=2, values_only=True):
            # Se a linha não tiver data ou for a nossa linha de "TOTAIS", ele ignora
            if not row[0] or str(row[0]).strip() == "TOTAIS":
                continue
            
            try:
                # O nosso gerador_excel colocou nesta ordem exata:
                # 0: Data, 1: Banco, 2: Categoria, 3: Descrição, 4: Tipo, 5: Valor
                data = str(row[0]).strip()
                banco = str(row[1]).strip() if row[1] else "Desconhecido"
                categoria = str(row[2]).strip() if row[2] else "Outras Despesas"
                descricao = str(row[3]).strip() if row[3] else ""
                
                # Normaliza o tipo (para o frontend entender fácil)
                tipo_raw = str(row[4]).strip().lower()
                tipo = "entrada" if "entrada" in tipo_raw else "saida"
                
                # Normaliza o valor (garantindo que é um número decimal)
                valor = float(row[5]) if row[5] is not None else 0.0

                # Adiciona no nosso "panelão" de consolidação
                transacoes_consolidadas.append({
                    "data": data,
                    "banco": banco,
                    "categoria": categoria,
                    "descricao": descricao,
                    "tipo": tipo,
                    "valor": valor
                })
            except Exception:
                # Se uma linha estiver completamente quebrada, ele ignora e continua
                continue

    return transacoes_consolidadas