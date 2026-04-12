"""
Gerador e leitor de planilhas-modelo para dashboards personalizados do Controllo.
Suporta o tipo 'Sazonal' — pode ser expandido para outros tipos futuramente.
"""
from __future__ import annotations
import io
import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MESES_PT = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]

# ──────────────────────────────────────────────────────────────────
#  Estilos reutilizáveis
# ──────────────────────────────────────────────────────────────────
def _estilos():
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill  = PatternFill("solid", start_color="1E3A5F", end_color="1E3A5F")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    normal_font  = Font(name="Calibri", size=10)
    title_font   = Font(name="Calibri", bold=True, size=14, color="1E3A5F")
    return border, header_font, header_fill, header_align, normal_font, title_font


# ──────────────────────────────────────────────────────────────────
#  Gerador do template (vazio ou com dados de demonstração)
# ──────────────────────────────────────────────────────────────────
def gerar_template_sazonal(com_demo: bool = False) -> bytes:
    """
    Gera um modelo Excel para dashboard sazonal.
    - com_demo=False → planilha em branco (o usuário preenche).
    - com_demo=True  → planilha pré-preenchida com dados fictícios realistas.
    """
    wb = Workbook()
    border, hf, hfill, ha, nf, tf = _estilos()

    # ── ABA 1: INSTRUÇÕES ─────────────────────────────────────────
    ws_inst = wb.active
    ws_inst.title = "Instruções"
    ws_inst.column_dimensions["A"].width = 90
    ws_inst.row_dimensions[1].height = 36

    instrucoes = [
        ("CONTROLLO — MODELO DE DASHBOARD SAZONAL", True,  14, "FFFFFF", "1E3A5F"),
        ("",                                         False, 10, "333333", None),
        ("Como preencher esta planilha:",            True,  12, "1E3A5F", None),
        ("",                                         False, 10, "333333", None),
        ("1.  Acesse a aba 'Dashboard_Sazonal'.",    False, 10, "333333", None),
        ("2.  Preencha as informações da empresa na aba 'Configurações'.", False, 10, "333333", None),
        ("3.  Insira os valores mensais nas colunas B até G (linhas 2 a 13).", False, 10, "333333", None),
        ("4.  NÃO altere os cabeçalhos, nomes de abas nem a estrutura da planilha.", True, 10, "C62828", None),
        ("5.  Salve o arquivo e faça o upload em 'Dashboard → Dashboard Sazonal'.", False, 10, "333333", None),
        ("",                                         False, 10, "333333", None),
        ("Colunas obrigatórias:",                    True,  11, "1E3A5F", None),
        ("  • Receita Bruta (R$)            — total de receitas/vendas do mês",      False, 10, "333333", None),
        ("  • Custos (R$)                   — custo dos produtos/serviços vendidos", False, 10, "333333", None),
        ("  • Despesas Operacionais (R$)    — despesas adm. + comerciais",            False, 10, "333333", None),
        ("  • Lucro Líquido (R$)            — resultado final do mês (pode ser negativo)", False, 10, "333333", None),
        ("  • Entradas de Caixa (R$)        — total de recebimentos no mês",         False, 10, "333333", None),
        ("  • Saídas de Caixa (R$)          — total de pagamentos no mês",           False, 10, "333333", None),
        ("",                                         False, 10, "333333", None),
        ("Dica: Células em branco são interpretadas como zero.",  False, 9, "666666", None),
    ]

    for i, (txt, bold, sz, cor, bg) in enumerate(instrucoes, 1):
        cell = ws_inst.cell(row=i, column=1, value=txt)
        cell.font = Font(name="Calibri", bold=bold, size=sz, color=cor)
        if bg:
            cell.fill = PatternFill("solid", start_color=bg, end_color=bg)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_inst.row_dimensions[i].height = 28 if i == 1 else 18

    # ── ABA 2: CONFIGURAÇÕES ──────────────────────────────────────
    ws_conf = wb.create_sheet("Configurações")
    ws_conf.column_dimensions["A"].width = 30
    ws_conf.column_dimensions["B"].width = 44

    campos_conf = [
        ("Campo",               "Valor"),
        ("Nome da Empresa",     "Minha Empresa LTDA" if com_demo else ""),
        ("Ano de Referência",   "2025"  if com_demo else str(datetime.date.today().year)),
        ("Setor",               "Varejo" if com_demo else ""),
        ("Tipo de Dashboard",   "Sazonal"),
    ]
    for i, (campo, valor) in enumerate(campos_conf, 1):
        ca = ws_conf.cell(row=i, column=1, value=campo)
        cb = ws_conf.cell(row=i, column=2, value=valor)
        ca.border = border
        cb.border = border
        if i == 1:
            for c in (ca, cb):
                c.font = hf; c.fill = hfill; c.alignment = ha
        else:
            ca.font = Font(name="Calibri", bold=True, size=10)
            cb.font = Font(name="Calibri", size=10)

    # ── ABA 3: DASHBOARD SAZONAL ──────────────────────────────────
    ws = wb.create_sheet("Dashboard_Sazonal")

    cabecalhos = [
        "Mês",
        "Receita Bruta (R$)",
        "Custos (R$)",
        "Despesas Operacionais (R$)",
        "Lucro Líquido (R$)",
        "Entradas de Caixa (R$)",
        "Saídas de Caixa (R$)",
    ]
    larguras = [16, 22, 18, 30, 20, 24, 20]

    for col, (cab, larg) in enumerate(zip(cabecalhos, larguras), 1):
        cell = ws.cell(row=1, column=col, value=cab)
        cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = larg
    ws.row_dimensions[1].height = 30

    # Dados de demonstração realistas (PME de varejo, sazonalidade de fim de ano)
    demo_rows = [
        (120_000, 45_000, 38_000,  37_000, 115_000,  90_000),
        (115_000, 43_000, 36_000,  36_000, 110_000,  88_000),
        (130_000, 49_000, 40_000,  41_000, 125_000,  95_000),
        (145_000, 54_000, 42_000,  49_000, 140_000, 100_000),
        (160_000, 60_000, 45_000,  55_000, 155_000, 108_000),
        (175_000, 65_000, 47_000,  63_000, 170_000, 115_000),
        (180_000, 67_000, 48_000,  65_000, 175_000, 118_000),
        (190_000, 71_000, 50_000,  69_000, 185_000, 125_000),
        (185_000, 69_000, 49_000,  67_000, 180_000, 122_000),
        (200_000, 75_000, 52_000,  73_000, 195_000, 130_000),
        (210_000, 78_000, 54_000,  78_000, 205_000, 135_000),
        (250_000, 93_000, 60_000,  97_000, 245_000, 155_000),
    ]

    row_fills = [
        PatternFill("solid", start_color="F0F4FF", end_color="F0F4FF"),
        PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF"),
    ]

    for idx, mes_nome in enumerate(MESES_PT, 1):
        row = idx + 1
        ws.row_dimensions[row].height = 20
        c_mes = ws.cell(row=row, column=1, value=mes_nome)
        c_mes.font = Font(name="Calibri", bold=True, size=10)
        c_mes.fill = row_fills[(idx - 1) % 2]
        c_mes.border = border

        valores = demo_rows[idx - 1] if com_demo else ("", "", "", "", "", "")
        for col, val in enumerate(valores, 2):
            cell = ws.cell(row=row, column=col, value=val if com_demo else None)
            cell.font = nf
            cell.fill = row_fills[(idx - 1) % 2]
            cell.border = border
            if com_demo and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ──────────────────────────────────────────────────────────────────
#  Leitor do template preenchido
# ──────────────────────────────────────────────────────────────────
def ler_template_sazonal(conteudo: bytes) -> dict:
    """
    Lê uma planilha de template sazonal preenchida e retorna os dados
    estruturados para renderização do dashboard.
    """
    wb = load_workbook(filename=io.BytesIO(conteudo), data_only=True)

    # Lê configurações
    nome_empresa = "Empresa"
    ano = datetime.date.today().year
    if "Configurações" in wb.sheetnames:
        for row in wb["Configurações"].iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                campo = str(row[0]).strip()
                valor = str(row[1]).strip()
                if campo == "Nome da Empresa" and valor:
                    nome_empresa = valor
                elif campo == "Ano de Referência":
                    try:
                        ano = int(valor)
                    except ValueError:
                        pass

    # Lê dados sazonais
    if "Dashboard_Sazonal" not in wb.sheetnames:
        raise ValueError(
            "Aba 'Dashboard_Sazonal' não encontrada. Certifique-se de usar o modelo correto."
        )

    ws = wb["Dashboard_Sazonal"]

    def _num(v) -> float:
        if v is None:
            return 0.0
        try:
            return float(v)
        except (TypeError, ValueError):
            return 0.0

    meses_data = []
    for row in ws.iter_rows(min_row=2, max_row=13, values_only=True):
        if not row[0]:
            continue
        meses_data.append({
            "mes":                   str(row[0]),
            "receita_bruta":         _num(row[1]),
            "custos":                _num(row[2]),
            "despesas_operacionais": _num(row[3]),
            "lucro_liquido":         _num(row[4]),
            "entradas_caixa":        _num(row[5]),
            "saidas_caixa":          _num(row[6]),
        })

    if not meses_data:
        raise ValueError(
            "Nenhum dado encontrado. Verifique se preencheu a aba 'Dashboard_Sazonal'."
        )

    # Totais
    total_receita = sum(m["receita_bruta"]         for m in meses_data)
    total_custos  = sum(m["custos"]                for m in meses_data)
    total_desp    = sum(m["despesas_operacionais"] for m in meses_data)
    total_lucro   = sum(m["lucro_liquido"]         for m in meses_data)
    total_entradas = sum(m["entradas_caixa"]       for m in meses_data)
    total_saidas   = sum(m["saidas_caixa"]         for m in meses_data)
    n = len(meses_data)

    pico  = max(meses_data, key=lambda m: m["receita_bruta"])
    vale  = min(meses_data, key=lambda m: m["receita_bruta"])

    return {
        "tipo":         "sazonal",
        "empresa":      nome_empresa,
        "ano":          ano,
        "meses":        meses_data,
        "totais": {
            "receita_bruta":         total_receita,
            "custos":                total_custos,
            "despesas_operacionais": total_desp,
            "lucro_liquido":         total_lucro,
            "entradas_caixa":        total_entradas,
            "saidas_caixa":          total_saidas,
        },
        "medias": {
            "receita_bruta":  total_receita / n if n else 0,
            "lucro_liquido":  total_lucro   / n if n else 0,
            "entradas_caixa": total_entradas / n if n else 0,
        },
        "pico_receita": pico["mes"],
        "vale_receita": vale["mes"],
        "margem_media": (total_lucro / total_receita * 100) if total_receita > 0 else 0,
    }
