"""
excel_reader.py — Lê planilhas Excel geradas pelo Leitor de Extrato do Controllo.

Preserva o número de linha Excel de cada transação para que os relatórios
de divergência possam apontar exatamente onde está o problema.

Formato esperado (gerado por gerador_excel.py):
  Coluna A: Data
  Coluna B: Banco
  Coluna C: Categoria
  Coluna D: Descrição
  Coluna E: Tipo  (Entrada / Saída)
  Coluna F: Valor (R$)
"""

from io import BytesIO
from openpyxl import load_workbook


def ler_excel_com_linhas(arquivos_bytes: list[bytes]) -> list[dict]:
    """
    Lê uma ou mais planilhas Excel e retorna todas as transações
    com o número de linha exato no arquivo Excel.

    Args:
        arquivos_bytes: lista de bytes de arquivos .xlsx

    Returns:
        lista de dicts com chaves:
            linha_excel, data, banco, categoria, descricao, tipo, valor
    """
    transacoes: list[dict] = []

    for idx_arquivo, arquivo_bytes in enumerate(arquivos_bytes):
        try:
            wb = load_workbook(filename=BytesIO(arquivo_bytes), data_only=True)
        except Exception as e:
            continue  # arquivo inválido — ignora

        nome_aba = "Transações" if "Transações" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[nome_aba]

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Pula linhas vazias, TOTAIS e separadores
            if not row[0]:
                continue
            val_a = str(row[0]).strip()
            if val_a in ("TOTAIS", "TOTAL", ""):
                continue
            # Ignora linha se parece ser separador (só texto sem data no padrão DD/MM/YYYY)
            import re as _re
            if not _re.match(r"\d{1,2}[/\-]\d{1,2}", val_a):
                continue

            try:
                data      = val_a
                banco     = str(row[1]).strip() if row[1] else "Desconhecido"
                categoria = str(row[2]).strip() if row[2] else ""
                descricao = str(row[3]).strip() if row[3] else ""
                tipo_raw  = str(row[4]).strip().lower() if row[4] else ""
                tipo      = "entrada" if "entrada" in tipo_raw else "saida"
                valor     = float(row[5]) if row[5] is not None else 0.0
            except Exception:
                continue

            if valor == 0.0:
                continue

            transacoes.append({
                "linha_excel": row_num,
                "data":        data,
                "banco":       banco,
                "categoria":   categoria,
                "descricao":   descricao,
                "tipo":        tipo,
                "valor":       abs(valor),  # sempre positivo
            })

    return transacoes
