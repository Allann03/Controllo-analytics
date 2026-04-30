"""
gerador_excel.py — Excel básico no formato contábil (Sessão 22, spec Allan).

Layout fixo de 7 colunas, 1 aba, sem totalizadores. Destinado a importação
em sistema contábil. Para o Excel contábil com partida dobrada e plano de
contas, ver gerador_excel_contabil.py (não tocado nesta sessão).
"""

from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


_CABECALHOS = ["Lançamento", "Data", "Débito", "Crédito", "Valor", "Histórico", "Complemento"]
_LARGURAS = [12, 12, 10, 10, 14, 10, 50]
_NUM_COLUNAS = len(_CABECALHOS)

# Estilos S22 — extensão visual (Sessão 22 ext.):
# cabeçalho azul sério com texto branco, linhas verde-claro (entrada) e rosa-claro (saída).
_HEADER_FILL = PatternFill("solid", start_color="FF1F4E79", end_color="FF1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_ROW_FILL_ENTRADA = PatternFill("solid", start_color="FFE2EFDA", end_color="FFE2EFDA")
_ROW_FILL_SAIDA = PatternFill("solid", start_color="FFFCE4D6", end_color="FFFCE4D6")

# Estilos S22 ext.2 — alinhamento + bordas (padrão Contmatic).
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
_THIN_SIDE = Side(style="thin", color="FF000000")
_BORDER_THIN = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

# Mapa coluna (1-indexada) -> alinhamento horizontal nas linhas de dados.
# A=Lançamento esquerda | B=Data centro | C=Débito centro | D=Crédito centro
# E=Valor direita | F=Histórico esquerda | G=Complemento esquerda
_ALINHAMENTO_DADOS = {
    1: _ALIGN_LEFT,
    2: _ALIGN_CENTER,
    3: _ALIGN_CENTER,
    4: _ALIGN_CENTER,
    5: _ALIGN_RIGHT,
    6: _ALIGN_LEFT,
    7: _ALIGN_LEFT,
}


def _to_float(val) -> float:
    if isinstance(val, Decimal):
        return float(val)
    if val is None:
        return 0.0
    return float(val)


def _parse_data(data_str):
    """'DD/MM/YYYY' -> datetime.date. Retorna o próprio str se não parsear."""
    if isinstance(data_str, datetime):
        return data_str.date()
    try:
        return datetime.strptime(str(data_str), "%d/%m/%Y").date()
    except (ValueError, TypeError):
        return data_str


def gerar_excel(transacoes: list, caminho_saida: str, saldo_inicial: float = None) -> str:
    """Gera Excel básico no layout S22 (7 colunas, 1 aba, sem totalizadores).

    saldo_inicial: aceito por compatibilidade com chamadas legadas. Não é mais
    usado no Excel básico (S22 — spec Allan). O Excel contábil
    (gerador_excel_contabil.py) ainda usa saldo_inicial.
    """
    # Filtra tipo='posicao' (saldos de carteira de investimento, não movimentação de caixa).
    # Decisão da Sessão 22 — spec do Allan é Excel contábil, posições não são lançamentos.
    transacoes_filtradas = [t for t in (transacoes or []) if t.get("tipo") != "posicao"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"

    # Cabeçalho: azul + branco + bold + centralizado, sem borda.
    for col_idx, (cab, larg) in enumerate(zip(_CABECALHOS, _LARGURAS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=cab)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _ALIGN_CENTER
        ws.column_dimensions[cell.column_letter].width = larg

    for i, t in enumerate(transacoes_filtradas, start=1):
        row = i + 1
        valor = abs(_to_float(t.get("valor", 0)))
        data_val = _parse_data(t.get("data", ""))
        fill_linha = _ROW_FILL_ENTRADA if t.get("tipo") == "entrada" else _ROW_FILL_SAIDA
        # Title Case na coluna Complemento (S22 ext.2). Allan aceitou que siglas
        # como PIX/TED/DOC fiquem como "Pix"/"Ted"/"Doc".
        complemento = str(t.get("descricao", "") or "").title()

        ws.cell(row=row, column=1, value=i)
        c_data = ws.cell(row=row, column=2, value=data_val)
        if isinstance(data_val, (datetime,)) or hasattr(data_val, "year"):
            c_data.number_format = "DD/MM/YYYY"
        c_valor = ws.cell(row=row, column=5, value=valor)
        c_valor.number_format = "#,##0.00"
        ws.cell(row=row, column=7, value=complemento)

        # Aplica fill, alinhamento e borda em A:G (linha inteira).
        for col_idx in range(1, _NUM_COLUNAS + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = fill_linha
            cell.alignment = _ALINHAMENTO_DADOS[col_idx]
            cell.border = _BORDER_THIN

    wb.save(caminho_saida)
    return caminho_saida
